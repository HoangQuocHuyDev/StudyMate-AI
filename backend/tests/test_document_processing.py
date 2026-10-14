"""Unit tests for the transactional document processing pipeline."""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import fields
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)
from sqlalchemy import (
    BigInteger,
    Boolean,
    Column,
    DateTime,
    MetaData,
    String,
    Table,
    Text,
    create_engine,
    event,
    inspect,
    select,
)
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import app.services.document_processing as processing
from app.models.document import Document
from app.models.document_deletion_job import DocumentDeletionJob
from app.models.document_chunk import DocumentChunk
from app.models.user import User
from app.services.document_chunk_repository import (
    DocumentChunkPersistenceError,
    count_document_chunks,
    get_document_chunks,
)
from app.services.document_extractor import extract_document_text
from app.services.document_processing import (
    DocumentNotFoundError,
    DocumentNotProcessableError,
    DocumentProcessingFailedError,
    DocumentProcessingInfrastructureError,
    DocumentProcessingResult,
    process_document,
    retry_document,
)
from app.services.document_service import create_document
from app.services.file_storage import FileStorage
from app.services.text_chunker import TextChunkingError


@pytest.fixture
def processing_environment(
    tmp_path: Path,
) -> Iterator[tuple[sessionmaker[Session], FileStorage, object]]:
    """Provide isolated SQLite metadata and temporary private storage."""

    database_engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(database_engine, "connect")
    def enable_sqlite_foreign_keys(dbapi_connection: object, record: object) -> None:
        cursor = dbapi_connection.cursor()  # type: ignore[attr-defined]
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    metadata = MetaData()
    Table(
        "users",
        metadata,
        Column("id", BigInteger, primary_key=True),
        Column("full_name", String(100), nullable=False),
        Column("username", String(30), nullable=False),
        Column("password_hash", Text, nullable=False),
        Column("role", String(10), nullable=False),
        Column("is_active", Boolean, nullable=False),
        Column("token_version", BigInteger, nullable=False),
        Column("created_at", DateTime(timezone=True), nullable=False),
    )
    Document.__table__.to_metadata(metadata)
    DocumentDeletionJob.__table__.to_metadata(metadata)
    DocumentChunk.__table__.to_metadata(metadata)
    metadata.create_all(database_engine)
    factory = sessionmaker(
        bind=database_engine,
        autoflush=False,
        expire_on_commit=False,
    )
    storage = FileStorage(
        tmp_path / "processing-documents",
        max_file_size=2 * 1024 * 1024,
    )
    yield factory, storage, database_engine
    database_engine.dispose()


def make_docx(text: str) -> bytes:
    document = DocxDocument()
    if text:
        document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf(
    page_texts: tuple[str | None, ...],
    *,
    encrypted: bool = False,
) -> bytes:
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        if text is None:
            continue
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_reference = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        content = DecodedStreamObject()
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content.set_data(
            f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode("ascii")
        )
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("test-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def add_document(
    factory: sessionmaker[Session],
    storage: FileStorage,
    *,
    document_id: int = 1,
    user_id: int = 7,
    file_type: str = "DOCX",
    content: bytes | None = None,
    status: str = "UPLOADED",
    store_file: bool = True,
) -> str:
    if content is None:
        content = make_docx("StudyMate lesson content")
    if store_file:
        storage_key = storage.save(io.BytesIO(content), user_id=user_id).storage_key
    else:
        storage_key = f"users/{user_id}/missing-document"
    now = datetime.now(timezone.utc)
    with factory() as session:
        if session.get(User, user_id) is None:
            session.add(
                User(
                    id=user_id,
                    full_name=f"User {user_id}",
                    username=f"user_{user_id}",
                    password_hash="test-only-hash",
                    role="USER",
                    is_active=True,
                    created_at=now,
                )
            )
            session.flush()
        session.add(
            Document(
                id=document_id,
                user_id=user_id,
                title=f"Document {document_id}",
                original_filename=f"document-{document_id}.{file_type.lower()}",
                storage_key=storage_key,
                file_type=file_type,
                mime_type={
                    "PDF": "application/pdf",
                    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                }[file_type],
                file_size=max(1, len(content)),
                status=status,
                created_at=now,
                updated_at=now,
            )
        )
        session.commit()
    return storage_key


def document_status(factory: sessionmaker[Session], document_id: int = 1) -> str:
    with factory() as session:
        return session.scalar(
            select(Document.status).where(Document.id == document_id)
        )


def document_processing_metadata(
    factory: sessionmaker[Session],
    document_id: int = 1,
) -> Document:
    with factory() as session:
        document = session.get(Document, document_id)
        assert document is not None
        session.expunge(document)
        return document


def read_stored_bytes(storage: FileStorage, storage_key: str) -> bytes:
    with storage.open_binary(storage_key) as stream:
        return stream.read()


def install_commit_failure(session: Session, commit_number: int) -> None:
    attempts = 0

    @event.listens_for(session, "before_commit")
    def fail_selected_commit(current_session: Session) -> None:
        nonlocal attempts
        attempts += 1
        if attempts == commit_number:
            raise SQLAlchemyError("test-only commit failure")


def test_uploaded_transitions_through_processing_to_ready(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)
    observed_statuses: list[str] = []
    real_extract = processing.extract_document_text

    def observe_processing(*args: object, **kwargs: object):
        with factory() as observer:
            observed_statuses.append(
                observer.scalar(select(Document.status).where(Document.id == 1))
            )
        return real_extract(*args, **kwargs)

    monkeypatch.setattr(processing, "extract_document_text", observe_processing)
    with factory() as session:
        result = process_document(session, storage, document_id=1, user_id=7)

    assert observed_statuses == ["PROCESSING"]
    assert result.status == "READY"
    assert document_status(factory) == "READY"


def test_processing_failure_transitions_to_failed(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, content=b"  \r\n\t  ")

    with factory() as session, pytest.raises(DocumentProcessingFailedError):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"


def test_missing_document_uses_not_found_error(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment

    with factory() as session, pytest.raises(DocumentNotFoundError):
        process_document(session, storage, document_id=999, user_id=7)


def test_other_users_document_is_indistinguishable_from_missing(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, user_id=8)

    with factory() as session, pytest.raises(DocumentNotFoundError):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "UPLOADED"


@pytest.mark.parametrize("current_status", ["READY", "FAILED", "PROCESSING"])
def test_non_uploaded_document_cannot_be_reprocessed(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    current_status: str,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, status=current_status)

    with factory() as session, pytest.raises(DocumentNotProcessableError):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == current_status


def test_atomic_claim_prevents_a_second_worker(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)
    real_extract = processing.extract_document_text
    competing_attempts = 0

    def compete_after_claim(*args: object, **kwargs: object):
        nonlocal competing_attempts
        competing_attempts += 1
        with factory() as competing_session:
            with pytest.raises(DocumentNotProcessableError):
                process_document(
                    competing_session,
                    storage,
                    document_id=1,
                    user_id=7,
                )
        return real_extract(*args, **kwargs)

    monkeypatch.setattr(processing, "extract_document_text", compete_after_claim)
    with factory() as first_session:
        result = process_document(
            first_session,
            storage,
            document_id=1,
            user_id=7,
        )

    assert competing_attempts == 1
    assert result.status == "READY"


def test_retry_failed_document_succeeds_and_updates_metadata(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, status="FAILED")

    with factory() as session:
        result = retry_document(session, storage, document_id=1, user_id=7)

    document = document_processing_metadata(factory)
    assert result.status == "READY"
    assert document.status == "READY"
    assert document.character_count == result.character_count > 0
    assert document.page_count is None
    assert document.processed_at is not None
    assert document.rag_status == "PENDING"
    assert document.indexed_at is None
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == result.chunk_count


def test_failed_retry_failure_keeps_failed_without_partial_data(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(
        factory,
        storage,
        status="FAILED",
        content=b"invalid utf-8: \xff\xfe",
    )

    with factory() as session, pytest.raises(DocumentProcessingFailedError):
        retry_document(session, storage, document_id=1, user_id=7)

    document = document_processing_metadata(factory)
    assert document.status == "FAILED"
    assert document.character_count is None
    assert document.page_count is None
    assert document.processed_at is None
    assert document.rag_status == "PENDING"
    assert document.indexed_at is None
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == 0


def test_failed_retry_commit_failure_rolls_back_chunks_and_metadata(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, status="FAILED")

    with factory() as session:
        install_commit_failure(session, 2)
        with pytest.raises(DocumentProcessingInfrastructureError):
            retry_document(session, storage, document_id=1, user_id=7)

    document = document_processing_metadata(factory)
    assert document.status == "PROCESSING"
    assert document.character_count is None
    assert document.page_count is None
    assert document.processed_at is None
    assert document.rag_status == "PENDING"
    assert document.indexed_at is None
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == 0


def test_chunkless_ready_document_can_be_reprocessed(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, status="READY")

    with factory() as session:
        result = retry_document(session, storage, document_id=1, user_id=7)

    assert result.status == "READY"
    document = document_processing_metadata(factory)
    assert document.character_count == result.character_count
    assert document.processed_at is not None
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == result.chunk_count


def test_ready_document_with_chunks_cannot_be_reprocessed_or_duplicated(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)
    with factory() as session:
        first_result = process_document(session, storage, document_id=1, user_id=7)

    with factory() as session, pytest.raises(DocumentNotProcessableError):
        retry_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "READY"
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == first_result.chunk_count


def test_failed_document_with_chunks_cannot_be_retried_or_duplicated(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)
    with factory() as session:
        first_result = process_document(session, storage, document_id=1, user_id=7)
    with factory() as session:
        document = session.get(Document, 1)
        assert document is not None
        document.status = "FAILED"
        session.commit()

    with factory() as session, pytest.raises(DocumentNotProcessableError):
        retry_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == first_result.chunk_count


def test_processing_document_is_not_claimed_for_recovery(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, status="PROCESSING")

    with factory() as session, pytest.raises(DocumentNotProcessableError):
        retry_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "PROCESSING"


def test_atomic_retry_claim_prevents_a_second_worker(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, status="FAILED")
    real_extract = processing.extract_document_text
    competing_attempts = 0

    def compete_after_retry_claim(*args: object, **kwargs: object):
        nonlocal competing_attempts
        competing_attempts += 1
        with factory() as competing_session:
            with pytest.raises(DocumentNotProcessableError):
                retry_document(
                    competing_session,
                    storage,
                    document_id=1,
                    user_id=7,
                )
        return real_extract(*args, **kwargs)

    monkeypatch.setattr(
        processing,
        "extract_document_text",
        compete_after_retry_claim,
    )
    with factory() as session:
        result = retry_document(session, storage, document_id=1, user_id=7)

    assert competing_attempts == 1
    assert result.status == "READY"


def test_user_cannot_retry_another_users_document(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, user_id=8, status="FAILED")

    with factory() as session, pytest.raises(DocumentNotFoundError):
        retry_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"


def test_extraction_runs_outside_a_database_transaction(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)
    real_extract = processing.extract_document_text
    transaction_states: list[bool] = []

    with factory() as session:
        def observe_transaction(*args: object, **kwargs: object):
            transaction_states.append(session.in_transaction())
            return real_extract(*args, **kwargs)

        monkeypatch.setattr(
            processing,
            "extract_document_text",
            observe_transaction,
        )
        process_document(session, storage, document_id=1, user_id=7)

    assert transaction_states == [False]


@pytest.mark.parametrize(
    ("file_type", "content", "expected_page_count"),
    [
        ("PDF", make_pdf(("StudyMate PDF lesson",)), 1),
        ("DOCX", make_docx("StudyMate DOCX lesson"), None),
    ],
    ids=("pdf", "docx"),
)
def test_supported_document_types_process_successfully(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    file_type: str,
    content: bytes,
    expected_page_count: int | None,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, file_type=file_type, content=content)
    started_at = datetime.now(timezone.utc)

    with factory() as session:
        result = process_document(session, storage, document_id=1, user_id=7)

    document = document_processing_metadata(factory)
    processed_at = document.processed_at
    assert processed_at is not None
    if processed_at.tzinfo is None:
        processed_at = processed_at.replace(tzinfo=timezone.utc)
    assert result.status == "READY"
    assert result.character_count > 0
    assert result.chunk_count > 0
    assert result.page_count == expected_page_count
    assert document.status == "READY"
    assert document.character_count == result.character_count
    assert document.page_count == expected_page_count
    assert started_at.replace(microsecond=0) <= processed_at <= datetime.now(timezone.utc)
    assert document.rag_status == "PENDING"
    assert document.indexed_at is None


def test_missing_stored_file_marks_document_failed_without_leaking_key(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    key = add_document(factory, storage, store_file=False)

    with factory() as session, pytest.raises(DocumentProcessingFailedError) as error:
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"
    assert key not in str(error.value)


@pytest.mark.parametrize(
    ("file_type", "content"),
    [
        ("DOCX", b"not a docx archive"),
        ("PDF", make_pdf(("Encrypted",), encrypted=True)),
        ("DOCX", make_docx("")),
    ],
    ids=("damaged-docx", "encrypted-pdf", "empty-docx"),
)
def test_content_failures_mark_document_failed(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    file_type: str,
    content: bytes,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, file_type=file_type, content=content)

    with factory() as session, pytest.raises(DocumentProcessingFailedError):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"


def test_chunking_error_marks_document_failed(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)

    def fail_chunking(*args: object, **kwargs: object) -> None:
        raise TextChunkingError("test-only chunk failure")

    monkeypatch.setattr(processing, "chunk_extracted_text", fail_chunking)
    with factory() as session, pytest.raises(DocumentProcessingFailedError):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"


def test_empty_chunk_output_is_rejected_and_marked_failed(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)
    monkeypatch.setattr(processing, "chunk_extracted_text", lambda extracted: [])

    with factory() as session, pytest.raises(DocumentProcessingFailedError):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "FAILED"


def test_chunk_insert_error_rolls_back_and_never_marks_ready(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)

    def fail_insert(*args: object, **kwargs: object) -> None:
        raise DocumentChunkPersistenceError("test-only insert failure")

    monkeypatch.setattr(processing, "save_document_chunks", fail_insert)
    with factory() as session, pytest.raises(
        DocumentProcessingInfrastructureError
    ):
        process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "PROCESSING"
    document = document_processing_metadata(factory)
    assert document.character_count is None
    assert document.page_count is None
    assert document.processed_at is None
    assert document.rag_status == "PENDING"
    assert document.indexed_at is None
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == 0


def test_result_has_valid_chunk_count_offsets_and_metadata(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    text = ("StudyMate artificial intelligence lesson. " * 30).strip()
    key = add_document(factory, storage, content=make_docx(text))

    with factory() as session:
        result = process_document(session, storage, document_id=1, user_id=7)
    extracted = extract_document_text(storage, storage_key=key, file_type="DOCX")

    assert result.chunk_count == len(result.chunks) > 1
    assert [chunk.index for chunk in result.chunks] == list(
        range(result.chunk_count)
    )
    for chunk in result.chunks:
        assert extracted.text[chunk.start_char : chunk.end_char] == chunk.text
        assert chunk.character_count == len(chunk.text) <= 800
        assert chunk.page_number is None


def test_processing_error_never_deletes_original_file(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    original = b"invalid utf-8: \xff\xfe"
    key = add_document(factory, storage, content=original)

    with factory() as session, pytest.raises(DocumentProcessingFailedError):
        process_document(session, storage, document_id=1, user_id=7)

    assert read_stored_bytes(storage, key) == original


def test_chunks_and_metadata_are_persisted_before_document_becomes_ready(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    factory, storage, database_engine = processing_environment
    add_document(factory, storage, content=make_docx("content " * 200))
    real_save = processing.save_document_chunks
    observed_finalize_state: list[tuple[str, int | None, datetime | None, int]] = []

    def observe_save(*args: object, **kwargs: object) -> int:
        inserted = real_save(*args, **kwargs)
        session = args[0]
        document = session.get(Document, 1)
        assert document is not None
        observed_finalize_state.append(
            (
                document.status,
                document.character_count,
                document.processed_at,
                inserted,
            )
        )
        return inserted

    monkeypatch.setattr(processing, "save_document_chunks", observe_save)
    with factory() as session:
        result = process_document(session, storage, document_id=1, user_id=7)

    assert result.chunk_count > 1
    assert observed_finalize_state == [
        ("PROCESSING", None, None, result.chunk_count)
    ]
    assert "document_chunks" in set(inspect(database_engine).get_table_names())
    with factory() as session:
        stored_chunks = get_document_chunks(session, document_id=1, user_id=7)
        assert count_document_chunks(session, document_id=1, user_id=7) == result.chunk_count
        document = session.get(Document, 1)
        assert document is not None
        assert document.status == "READY"
        assert document.character_count == result.character_count
        assert document.processed_at is not None
    assert [chunk.chunk_index for chunk in stored_chunks] == list(
        range(result.chunk_count)
    )
    assert [chunk.content for chunk in stored_chunks] == [
        chunk.text for chunk in result.chunks
    ]


def test_claim_commit_failure_rolls_back_and_reports_infrastructure_error(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)

    with factory() as session:
        install_commit_failure(session, 1)
        with pytest.raises(DocumentProcessingInfrastructureError):
            process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "UPLOADED"


def test_ready_commit_failure_never_reports_success(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)

    with factory() as session:
        install_commit_failure(session, 2)
        with pytest.raises(DocumentProcessingInfrastructureError):
            process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "PROCESSING"
    document = document_processing_metadata(factory)
    assert document.character_count is None
    assert document.page_count is None
    assert document.processed_at is None
    assert document.rag_status == "PENDING"
    assert document.indexed_at is None
    with factory() as session:
        assert count_document_chunks(session, document_id=1, user_id=7) == 0


def test_failed_commit_failure_does_not_claim_failed_was_saved(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage, content=b"invalid utf-8: \xff\xfe")

    with factory() as session:
        install_commit_failure(session, 2)
        with pytest.raises(DocumentProcessingInfrastructureError):
            process_document(session, storage, document_id=1, user_id=7)

    assert document_status(factory) == "PROCESSING"


def test_public_result_never_contains_private_storage_metadata(
    processing_environment: tuple[sessionmaker[Session], FileStorage, object],
) -> None:
    factory, storage, _ = processing_environment
    add_document(factory, storage)

    with factory() as session:
        result = process_document(session, storage, document_id=1, user_id=7)

    assert isinstance(result, DocumentProcessingResult)
    field_names = {field.name for field in fields(result)}
    assert field_names == {
        "document_id",
        "status",
        "character_count",
        "chunk_count",
        "page_count",
        "chunks",
    }
    assert "storage_key" not in repr(result)
    assert "processing-documents" not in repr(result)


def test_upload_still_only_creates_an_uploaded_document(
    tmp_path: Path,
    fake_session_factory: type,
) -> None:
    storage = FileStorage(tmp_path / "upload-check", max_file_size=1024)
    session = fake_session_factory()

    document = create_document(
        session,
        storage,
        io.BytesIO(make_pdf(("StudyMate upload remains unchanged",))),
        filename="lesson.pdf",
        content_type="application/pdf",
        title="Lesson",
        user_id=7,
    )

    assert document.status == "UPLOADED"


def test_upload_does_not_automatically_activate_processing(
    tmp_path: Path,
    fake_session_factory: type,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    storage = FileStorage(tmp_path / "no-auto-processing", max_file_size=1024)
    session = fake_session_factory()

    def unexpected_processing(*args: object, **kwargs: object) -> None:
        raise AssertionError("Upload must not automatically process documents.")

    monkeypatch.setattr(processing, "process_document", unexpected_processing)
    document = create_document(
        session,
        storage,
        io.BytesIO(make_pdf(("Upload without automatic processing",))),
        filename="lesson.pdf",
        content_type="application/pdf",
        title=None,
        user_id=7,
    )

    assert document.status == "UPLOADED"
