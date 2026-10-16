"""API tests for owner-scoped document details, viewing, and downloads."""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document as DocxFile
from fastapi.testclient import TestClient
from sqlalchemy import (
    BigInteger,
    Boolean,
    Column,
    DateTime,
    MetaData,
    String,
    Table,
    Text,
    create_engine,
)
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.security import create_access_token
from app.db.session import get_db
from app.main import app
from app.models.document import Document
from app.models.document_deletion_job import DocumentDeletionJob
from app.models.user import User
from app.routers import documents as documents_router
from app.routers.documents import get_file_storage
from app.services.file_storage import FileStorage
from app.services.document_preview import (
    DocumentPreviewTooLargeError,
    get_owned_docx_preview,
)


PDF_BYTES = b"%PDF-1.7\nStudyMate test PDF\n%%EOF"
DOCX_BYTES = b"PK\x03\x04StudyMate test DOCX"


def build_docx_bytes(
    *paragraphs: str,
    table_rows: tuple[tuple[str, ...], ...] = (),
) -> bytes:
    document = DocxFile()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(
            rows=len(table_rows),
            cols=len(table_rows[0]),
        )
        for row_index, row_values in enumerate(table_rows):
            for column_index, value in enumerate(row_values):
                table.cell(row_index, column_index).text = value
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def headers(user_id: int) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(user_id, 0)}"}


@pytest.fixture
def document_files_environment(
    client: TestClient,
    tmp_path: Path,
) -> Iterator[tuple[TestClient, sessionmaker[Session], FileStorage]]:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    metadata = MetaData()
    Table(
        "users",
        metadata,
        Column("id", BigInteger, primary_key=True),
        Column("full_name", String(100), nullable=False),
        Column("username", String(30), nullable=False),
        Column("password_hash", Text, nullable=False),
        Column("role", String(10), nullable=False),
        Column("is_active", Boolean, nullable=False),
        Column("token_version", BigInteger, nullable=False),
        Column("created_at", DateTime(timezone=True), nullable=False),
    )
    Document.__table__.to_metadata(metadata)
    DocumentDeletionJob.__table__.to_metadata(metadata)
    metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    storage = FileStorage(tmp_path / "documents", max_file_size=10 * 1024 * 1024)

    def override_database() -> Iterator[Session]:
        with factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_database
    app.dependency_overrides[get_file_storage] = lambda: storage
    yield client, factory, storage
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_file_storage, None)
    engine.dispose()


def add_user(factory: sessionmaker[Session], user_id: int) -> None:
    with factory() as session:
        session.add(
            User(
                id=user_id,
                full_name=f"User {user_id}",
                username=f"user_{user_id}",
                password_hash="test-only-hash",
                role="USER",
                is_active=True,
                token_version=0,
                created_at=datetime.now(timezone.utc),
            )
        )
        session.commit()


def add_document(
    factory: sessionmaker[Session],
    storage: FileStorage,
    *,
    document_id: int,
    user_id: int,
    file_type: str,
    filename: str,
    content: bytes,
) -> str:
    add_user(factory, user_id)
    stored = storage.save(io.BytesIO(content), user_id=user_id)
    mime_type = {
        "PDF": "application/pdf",
        "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }[file_type]
    now = datetime.now(timezone.utc)
    with factory() as session:
        session.add(
            Document(
                id=document_id,
                user_id=user_id,
                title=f"Document {document_id}",
                original_filename=filename,
                storage_key=stored.storage_key,
                file_type=file_type,
                mime_type=mime_type,
                file_size=len(content),
                status="READY",
                rag_status="PENDING",
                created_at=now,
                updated_at=now,
            )
        )
        session.commit()
    return stored.storage_key


@pytest.mark.parametrize(
    "path",
    [
        "/documents/1",
        "/documents/1/view",
        "/documents/1/preview",
        "/documents/1/download",
    ],
)
def test_document_file_endpoints_require_valid_jwt(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    path: str,
) -> None:
    client, _, _ = document_files_environment
    assert client.get(path).status_code == 401
    assert client.get(path, headers={"Authorization": "Bearer invalid"}).status_code == 401


@pytest.mark.parametrize(
    "path",
    [
        "/documents/1",
        "/documents/1/view",
        "/documents/1/preview",
        "/documents/1/download",
    ],
)
def test_other_owner_and_missing_document_are_both_not_found(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    path: str,
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="PDF",
        filename="private.pdf",
        content=PDF_BYTES,
    )
    add_user(factory, 9)

    assert client.get(path, headers=headers(9)).status_code == 404
    assert client.get(path.replace("/1", "/999"), headers=headers(9)).status_code == 404


def test_document_detail_returns_safe_metadata(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    storage_key = add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="PDF",
        filename="bài giảng.pdf",
        content=PDF_BYTES,
    )

    response = client.get("/documents/1", headers=headers(8))

    assert response.status_code == 200
    assert response.json()["original_filename"] == "bài giảng.pdf"
    assert response.json()["rag_status"] == "PENDING"
    assert "storage_key" not in response.text
    assert storage_key not in response.text
    assert "user_id" not in response.json()


def test_pdf_view_is_inline_and_nosniff(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="PDF",
        filename="bài giảng.pdf",
        content=PDF_BYTES,
    )

    response = client.get("/documents/1/view", headers=headers(8))

    assert response.status_code == 200
    assert response.content == PDF_BYTES
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline;")
    assert "filename*=UTF-8''b%C3%A0i%20gi%E1%BA%A3ng.pdf" in response.headers["content-disposition"]
    assert response.headers["x-content-type-options"] == "nosniff"


@pytest.mark.parametrize(
    ("file_type", "filename", "content", "mime_type"),
    [
        ("PDF", "lesson.pdf", PDF_BYTES, "application/pdf"),
        (
            "DOCX",
            "bài học.docx",
            DOCX_BYTES,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    ],
)
def test_pdf_and_docx_download_with_safe_attachment_headers(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    file_type: str,
    filename: str,
    content: bytes,
    mime_type: str,
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type=file_type,
        filename=filename,
        content=content,
    )

    response = client.get("/documents/1/download", headers=headers(8))

    assert response.status_code == 200
    assert response.content == content
    assert response.headers["content-type"] == mime_type
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.headers["x-content-type-options"] == "nosniff"


def test_missing_stored_file_returns_safe_error(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    storage_key = add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="PDF",
        filename="missing.pdf",
        content=PDF_BYTES,
    )
    assert storage.delete(storage_key) is True

    response = client.get("/documents/1/download", headers=headers(8))

    assert response.status_code == 500
    assert storage_key not in response.text
    assert "storage" not in response.text.lower()


def test_docx_cannot_be_viewed_inline(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="lesson.docx",
        content=DOCX_BYTES,
    )

    response = client.get("/documents/1/view", headers=headers(8))

    assert response.status_code == 400


def test_ready_docx_preview_returns_normalized_plain_text(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    content = build_docx_bytes(
        "Tiêu đề bài học",
        "Đoạn văn thứ nhất.",
        table_rows=(("Khái niệm", "Nội dung"), ("AI", "Trí tuệ nhân tạo")),
    )
    storage_key = add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="bài học.docx",
        content=content,
    )

    response = client.get("/documents/1/preview", headers=headers(8))

    assert response.status_code == 200
    payload = response.json()
    assert payload == {
        "document_id": 1,
        "title": "Document 1",
        "original_filename": "bài học.docx",
        "file_type": "DOCX",
        "content": (
            "Tiêu đề bài học\n\nĐoạn văn thứ nhất.\n\n"
            "Khái niệm | Nội dung\n\nAI | Trí tuệ nhân tạo"
        ),
        "character_count": len(
            "Tiêu đề bài học\n\nĐoạn văn thứ nhất.\n\n"
            "Khái niệm | Nội dung\n\nAI | Trí tuệ nhân tạo"
        ),
    }
    assert storage_key not in response.text
    assert "user_id" not in payload


def test_empty_ready_docx_preview_returns_an_empty_content(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="empty.docx",
        content=build_docx_bytes(),
    )

    response = client.get("/documents/1/preview", headers=headers(8))

    assert response.status_code == 200
    assert response.json()["content"] == ""
    assert response.json()["character_count"] == 0


@pytest.mark.parametrize(
    "document_status",
    ["UPLOADED", "PROCESSING", "READY", "FAILED"],
)
def test_docx_preview_reads_the_source_file_in_every_processing_state(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    document_status: str,
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="pending.docx",
        content=build_docx_bytes("Nội dung"),
    )
    with factory() as session:
        document = session.get(Document, 1)
        assert document is not None
        document.status = document_status
        session.commit()

    response = client.get("/documents/1/preview", headers=headers(8))

    assert response.status_code == 200
    assert response.json()["content"] == "N\u1ed9i dung"
    with factory() as session:
        document = session.get(Document, 1)
        assert document is not None
        assert document.status == document_status


def test_tombstoned_docx_preview_is_not_found(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    storage_key = add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="deleting.docx",
        content=build_docx_bytes("Nội dung"),
    )
    with factory() as session:
        session.add(
            DocumentDeletionJob(
                document_id=1,
                user_id=8,
                storage_key=storage_key,
                status="PENDING",
                current_step="VECTOR_CLEANUP",
                attempt_count=0,
            )
        )
        session.commit()

    assert client.get("/documents/1/preview", headers=headers(8)).status_code == 404


def test_pdf_is_rejected_by_docx_preview_endpoint(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="PDF",
        filename="lesson.pdf",
        content=PDF_BYTES,
    )

    response = client.get("/documents/1/preview", headers=headers(8))

    assert response.status_code == 400


@pytest.mark.parametrize("missing", [False, True])
def test_missing_or_damaged_docx_preview_returns_safe_error(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    missing: bool,
) -> None:
    client, factory, storage = document_files_environment
    storage_key = add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="unavailable.docx",
        content=DOCX_BYTES,
    )
    if missing:
        assert storage.delete(storage_key) is True

    response = client.get("/documents/1/preview", headers=headers(8))

    assert response.status_code == 500
    assert storage_key not in response.text
    assert "storage" not in response.text.lower()


def test_docx_preview_enforces_the_existing_extraction_limit(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    _, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="long.docx",
        content=build_docx_bytes("Nội dung dài hơn giới hạn thử nghiệm."),
    )

    with factory() as session, pytest.raises(DocumentPreviewTooLargeError):
        get_owned_docx_preview(
            session,
            storage,
            document_id=1,
            user_id=8,
            max_characters=10,
        )


def test_docx_preview_limit_error_returns_413(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client, factory, storage = document_files_environment
    add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="large.docx",
        content=build_docx_bytes("N\u1ed9i dung"),
    )

    def reject_oversized_preview(*args: object, **kwargs: object) -> None:
        raise DocumentPreviewTooLargeError("test-only safe limit failure")

    monkeypatch.setattr(
        documents_router,
        "get_owned_docx_preview",
        reject_oversized_preview,
    )

    response = client.get("/documents/1/preview", headers=headers(8))

    assert response.status_code == 413
    assert "test-only" not in response.text


def test_docx_preview_does_not_modify_document_or_source_file(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    source = build_docx_bytes("Nội dung chỉ đọc")
    storage_key = add_document(
        factory,
        storage,
        document_id=1,
        user_id=8,
        file_type="DOCX",
        filename="readonly.docx",
        content=source,
    )
    with factory() as session:
        before = session.get(Document, 1)
        assert before is not None
        before_values = (before.status, before.updated_at, before.storage_key)

    assert client.get("/documents/1/preview", headers=headers(8)).status_code == 200

    with factory() as session:
        after = session.get(Document, 1)
        assert after is not None
        assert (after.status, after.updated_at, after.storage_key) == before_values
    with storage.open_binary(storage_key) as stream:
        assert stream.read() == source


def test_legacy_txt_download_is_rejected_before_storage_read(
    document_files_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = document_files_environment
    add_user(factory, 8)
    stored = storage.save(io.BytesIO(b"legacy"), user_id=8)
    now = datetime.now(timezone.utc)
    with factory() as session:
        session.execute(
            Document.__table__.insert().values(
                id=1,
                user_id=8,
                title="Legacy",
                original_filename="legacy.txt",
                storage_key=stored.storage_key,
                file_type="TXT",
                mime_type="text/plain",
                file_size=6,
                status="READY",
                rag_status="PENDING",
                created_at=now,
                updated_at=now,
            )
        )
        session.commit()

    response = client.get("/documents/1/download", headers=headers(8))

    assert response.status_code == 400
    assert response.json()["detail"] == "Chỉ hỗ trợ tải xuống tài liệu PDF hoặc DOCX."
