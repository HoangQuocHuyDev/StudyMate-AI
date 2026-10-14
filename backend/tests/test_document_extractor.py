"""Unit tests for bounded document text extraction from private storage."""

from __future__ import annotations

import io
import os
import subprocess
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.document_extractor import (
    DocumentAccessError,
    DocumentExtractionError,
    DocumentExtractionLimitError,
    EmptyDocumentTextError,
    EncryptedDocumentError,
    UnsupportedDocumentTypeError,
    extract_document_text,
)
from app.services.file_storage import FileStorage


@pytest.fixture
def extraction_storage(tmp_path: Path) -> tuple[FileStorage, Path]:
    root = tmp_path / "documents"
    return FileStorage(root, max_file_size=2 * 1024 * 1024), root


def store_bytes(storage: FileStorage, content: bytes) -> str:
    return storage.save(io.BytesIO(content), user_id=7).storage_key


def stored_path(root: Path, storage_key: str) -> Path:
    return root.joinpath(*storage_key.split("/"))


def make_docx(
    *,
    paragraphs: tuple[str, ...] = (),
    table_rows: tuple[tuple[str, ...], ...] = (),
) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(
            rows=len(table_rows),
            cols=len(table_rows[0]),
        )
        for row_index, values in enumerate(table_rows):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf(
    page_texts: tuple[str | None, ...],
    *,
    encrypted: bool = False,
) -> bytes:
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        if text is None:
            continue

        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_reference = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        content = DecodedStreamObject()
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content.set_data(
            f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode("ascii")
        )
        page[NameObject("/Contents")] = writer._add_object(content)

    if encrypted:
        writer.encrypt("test-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_extract_docx_paragraphs(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(
        storage,
        make_docx(paragraphs=("Mở đầu", "Nội dung bài học")),
    )

    result = extract_document_text(storage, storage_key=key, file_type="DOCX")

    assert result.text == "Mở đầu\n\nNội dung bài học"
    assert result.page_count is None


def test_extract_docx_table(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(
        storage,
        make_docx(table_rows=(("Thuật ngữ", "Ý nghĩa"), ("AI", "Trí tuệ"))),
    )

    result = extract_document_text(storage, storage_key=key, file_type="DOCX")

    assert "Thuật ngữ | Ý nghĩa" in result.text
    assert "AI | Trí tuệ" in result.text


def test_extract_docx_preserves_paragraph_and_table_order(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    content = make_docx(
        paragraphs=("Trước bảng",),
        table_rows=(("Cột A", "Cột B"),),
    )
    document = DocxDocument(io.BytesIO(content))
    document.add_paragraph("Sau bảng")
    output = io.BytesIO()
    document.save(output)
    key = store_bytes(storage, output.getvalue())

    result = extract_document_text(storage, storage_key=key, file_type="DOCX")

    assert result.text.index("Trước bảng") < result.text.index("Cột A | Cột B")
    assert result.text.index("Cột A | Cột B") < result.text.index("Sau bảng")


def test_extract_docx_rejects_damaged_archive(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, b"not a docx archive")

    with pytest.raises(DocumentExtractionError):
        extract_document_text(storage, storage_key=key, file_type="DOCX")


def test_extract_docx_rejects_document_without_text(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_docx())

    with pytest.raises(EmptyDocumentTextError):
        extract_document_text(storage, storage_key=key, file_type="DOCX")


def test_extract_pdf_text_layer(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_pdf(("StudyMate PDF",)))

    result = extract_document_text(storage, storage_key=key, file_type="PDF")

    assert result.text == "StudyMate PDF"
    assert result.page_count == 1


def test_extract_pdf_multiple_pages_preserves_boundary(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_pdf(("First page", "Second page")))

    result = extract_document_text(storage, storage_key=key, file_type="PDF")

    assert result.page_count == 2
    assert "First page" in result.text
    assert "Second page" in result.text
    assert "\f" in result.text


def test_extract_pdf_without_text_layer_requires_ocr(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_pdf((None,)))

    with pytest.raises(EmptyDocumentTextError, match="OCR"):
        extract_document_text(storage, storage_key=key, file_type="PDF")


def test_extract_pdf_rejects_encrypted_document(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_pdf(("Secret",), encrypted=True))

    with pytest.raises(EncryptedDocumentError):
        extract_document_text(storage, storage_key=key, file_type="PDF")


def test_extract_pdf_rejects_damaged_document(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, b"%PDF-1.7\ninvalid")

    with pytest.raises(DocumentExtractionError):
        extract_document_text(storage, storage_key=key, file_type="PDF")


def test_extract_pdf_enforces_page_limit(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_pdf(("One", "Two")))

    with pytest.raises(DocumentExtractionLimitError):
        extract_document_text(
            storage,
            storage_key=key,
            file_type="PDF",
            max_pdf_pages=1,
        )


@pytest.mark.parametrize("file_type", ["TXT", "XLSX"])
def test_extract_rejects_unsupported_type(
    extraction_storage: tuple[FileStorage, Path],
    file_type: str,
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, b"content")

    with pytest.raises(UnsupportedDocumentTypeError):
        extract_document_text(storage, storage_key=key, file_type=file_type)


def test_extract_rejects_missing_stored_file_without_leaking_key(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    missing_key = "users/7/00000000000000000000000000000000"

    with pytest.raises(DocumentAccessError) as error:
        extract_document_text(
            storage,
            storage_key=missing_key,
            file_type="PDF",
        )

    assert missing_key not in str(error.value)


def test_extract_enforces_input_file_size_limit(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_docx(paragraphs=("larger than limit",)))

    with pytest.raises(DocumentExtractionLimitError):
        extract_document_text(
            storage,
            storage_key=key,
            file_type="DOCX",
            max_file_size=4,
        )


def test_extract_enforces_character_limit(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage
    key = store_bytes(storage, make_docx(paragraphs=("more than ten characters",)))

    with pytest.raises(DocumentExtractionLimitError):
        extract_document_text(
            storage,
            storage_key=key,
            file_type="DOCX",
            max_characters=10,
        )


def test_extract_rejects_path_outside_storage_root(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, _ = extraction_storage

    with pytest.raises(DocumentAccessError):
        extract_document_text(
            storage,
            storage_key="../outside.pdf",
            file_type="PDF",
        )


def test_extract_does_not_follow_symlink_outside_storage_root(
    extraction_storage: tuple[FileStorage, Path],
    tmp_path: Path,
) -> None:
    storage, root = extraction_storage
    outside_directory = tmp_path / "outside"
    outside_directory.mkdir()
    outside_file = outside_directory / "outside.pdf"
    outside_file.write_text("must not be read", encoding="utf-8")
    link_path = root / "escape"
    try:
        os.symlink(outside_directory, link_path, target_is_directory=True)
    except OSError as exc:
        if os.name != "nt":
            pytest.skip(f"Symbolic links are unavailable: {exc}")
        junction = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(link_path), str(outside_directory)],
            check=False,
            capture_output=True,
            text=True,
        )
        if junction.returncode != 0:
            pytest.skip("Symbolic links and Windows junctions are unavailable.")

    with pytest.raises(DocumentAccessError):
        extract_document_text(
            storage,
            storage_key="escape/outside.pdf",
            file_type="PDF",
        )
    assert outside_file.read_text(encoding="utf-8") == "must not be read"


def test_extraction_never_changes_original_file(
    extraction_storage: tuple[FileStorage, Path],
) -> None:
    storage, root = extraction_storage
    original = make_docx(paragraphs=("Nội dung nguyên bản",))
    key = store_bytes(storage, original)
    source = stored_path(root, key)

    extract_document_text(storage, storage_key=key, file_type="DOCX")

    assert source.read_bytes() == original
