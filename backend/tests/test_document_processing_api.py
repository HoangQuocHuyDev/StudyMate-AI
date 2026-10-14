"""Isolated API tests for synchronous document processing and status."""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import datetime, timezone
from pathlib import Path
from threading import Event, Thread

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import (
    BigInteger,
    Boolean,
    Column,
    DateTime,
    MetaData,
    String,
    Table,
    Text,
    create_engine,
    event,
    func,
    select,
)
from sqlalchemy.orm import Session, sessionmaker

import app.services.document_processing as processing
from app.core.security import create_access_token
from app.db.session import get_db
from app.main import app
from app.models.document import Document
from app.models.document_deletion_job import DocumentDeletionJob
from app.models.document_chunk import DocumentChunk
from app.models.user import User
from app.routers.documents import get_file_storage
from app.services.file_storage import FileStorage


def authorization_header(user_id: int) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(user_id, 0)}"}


@pytest.fixture
def processing_api_environment(
    client: TestClient,
    tmp_path: Path,
) -> Iterator[tuple[TestClient, sessionmaker[Session], FileStorage]]:
    """Route the API to isolated SQLite and temporary private storage."""

    database_engine = create_engine(
        f"sqlite+pysqlite:///{tmp_path / 'processing-api.sqlite3'}",
        connect_args={"check_same_thread": False, "timeout": 5},
    )

    @event.listens_for(database_engine, "connect")
    def enable_sqlite_foreign_keys(dbapi_connection: object, record: object) -> None:
        cursor = dbapi_connection.cursor()  # type: ignore[attr-defined]
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    metadata = MetaData()
    Table(
        "users",
        metadata,
        Column("id", BigInteger, primary_key=True),
        Column("full_name", String(100), nullable=False),
        Column("username", String(30), nullable=False),
        Column("password_hash", Text, nullable=False),
        Column("role", String(10), nullable=False),
        Column("is_active", Boolean, nullable=False),
        Column("token_version", BigInteger, nullable=False),
        Column("created_at", DateTime(timezone=True), nullable=False),
    )
    Document.__table__.to_metadata(metadata)
    DocumentDeletionJob.__table__.to_metadata(metadata)
    DocumentChunk.__table__.to_metadata(metadata)
    metadata.create_all(database_engine)
    factory = sessionmaker(
        bind=database_engine,
        autoflush=False,
        expire_on_commit=False,
    )
    storage = FileStorage(tmp_path / "documents", max_file_size=2 * 1024 * 1024)

    def override_database() -> Iterator[Session]:
        with factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_database
    app.dependency_overrides[get_file_storage] = lambda: storage
    yield client, factory, storage
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_file_storage, None)
    database_engine.dispose()


def add_user(factory: sessionmaker[Session], user_id: int) -> None:
    with factory() as session:
        session.add(
            User(
                id=user_id,
                full_name=f"User {user_id}",
                username=f"user_{user_id}",
                password_hash="test-only-hash",
                role="USER",
                is_active=True,
                created_at=datetime.now(timezone.utc),
            )
        )
        session.commit()


def make_docx(text: str) -> bytes:
    document = DocxDocument()
    if text:
        document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def add_document(
    factory: sessionmaker[Session],
    storage: FileStorage,
    *,
    user_id: int = 42,
    document_id: int = 1,
    status: str = "UPLOADED",
    content: bytes | None = None,
) -> str:
    if content is None:
        content = make_docx("StudyMate API processing content")
    add_user(factory, user_id)
    stored = storage.save(io.BytesIO(content), user_id=user_id)
    now = datetime.now(timezone.utc)
    with factory() as session:
        session.add(
            Document(
                id=document_id,
                user_id=user_id,
                title="Processing API document",
                original_filename="processing-api.docx",
                storage_key=stored.storage_key,
                file_type="DOCX",
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                file_size=len(content),
                status=status,
                rag_status="PENDING",
                created_at=now,
                updated_at=now,
            )
        )
        session.commit()
    return stored.storage_key


def chunk_count(factory: sessionmaker[Session], document_id: int = 1) -> int:
    with factory() as session:
        return session.scalar(
            select(func.count())
            .select_from(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        ) or 0


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("post", "/documents/1/process"),
        ("get", "/documents/1/status"),
    ],
)
def test_processing_endpoints_require_jwt(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    method: str,
    path: str,
) -> None:
    client, _, _ = processing_api_environment

    response = getattr(client, method)(path)

    assert response.status_code == 401


def test_other_user_cannot_process_or_view_status(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage, user_id=8)
    add_user(factory, 42)
    headers = authorization_header(42)

    process_response = client.post("/documents/1/process", headers=headers)
    status_response = client.get("/documents/1/status", headers=headers)

    assert process_response.status_code == 404
    assert status_response.status_code == 404


def test_uploaded_document_processes_successfully(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage)

    response = client.post(
        "/documents/1/process",
        headers=authorization_header(42),
    )

    assert response.status_code == 200
    body = response.json()
    assert body["document_id"] == 1
    assert body["status"] == "READY"
    assert body["character_count"] > 0
    assert body["processed_at"] is not None
    assert body["chunk_count"] > 0
    assert body["rag_status"] == "PENDING"
    assert body["indexed_at"] is None


@pytest.mark.parametrize("initial_status", ["FAILED", "READY"])
def test_chunkless_retryable_document_processes_successfully(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    initial_status: str,
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage, status=initial_status)

    response = client.post(
        "/documents/1/process",
        headers=authorization_header(42),
    )

    assert response.status_code == 200
    assert response.json()["status"] == "READY"
    assert response.json()["chunk_count"] > 0


def test_ready_document_with_chunks_returns_conflict_without_duplicates(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage)
    headers = authorization_header(42)
    first_response = client.post("/documents/1/process", headers=headers)
    original_chunk_count = chunk_count(factory)

    retry_response = client.post("/documents/1/process", headers=headers)

    assert first_response.status_code == 200
    assert retry_response.status_code == 409
    assert chunk_count(factory) == original_chunk_count > 0


def test_processing_document_returns_conflict(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage, status="PROCESSING")

    response = client.post(
        "/documents/1/process",
        headers=authorization_header(42),
    )

    assert response.status_code == 409


def test_concurrent_requests_cannot_both_process_document(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage)
    headers = authorization_header(42)
    extraction_started = Event()
    allow_extraction = Event()
    real_extract = processing.extract_document_text
    responses: list[object] = []

    def blocking_extract(*args: object, **kwargs: object):
        extraction_started.set()
        assert allow_extraction.wait(timeout=5)
        return real_extract(*args, **kwargs)

    monkeypatch.setattr(processing, "extract_document_text", blocking_extract)

    def first_request() -> None:
        responses.append(client.post("/documents/1/process", headers=headers))

    worker = Thread(target=first_request)
    worker.start()
    assert extraction_started.wait(timeout=5)
    competing_response = client.post("/documents/1/process", headers=headers)
    allow_extraction.set()
    worker.join(timeout=5)

    assert worker.is_alive() is False
    assert competing_response.status_code == 409
    assert len(responses) == 1
    assert responses[0].status_code == 200  # type: ignore[attr-defined]
    assert chunk_count(factory) > 0


def test_status_returns_metadata_and_database_chunk_count(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = processing_api_environment
    add_document(factory, storage)
    headers = authorization_header(42)
    process_response = client.post("/documents/1/process", headers=headers)
    assert process_response.status_code == 200

    response = client.get("/documents/1/status", headers=headers)

    assert response.status_code == 200
    body = response.json()
    assert set(body) == {
        "document_id",
        "status",
        "rag_status",
        "character_count",
        "page_count",
        "processed_at",
        "indexed_at",
        "chunk_count",
    }
    assert body["document_id"] == 1
    assert body["status"] == "READY"
    assert body["rag_status"] == "PENDING"
    assert body["chunk_count"] == chunk_count(factory)


def test_processing_failure_returns_safe_response_without_internal_metadata(
    processing_api_environment: tuple[TestClient, sessionmaker[Session], FileStorage],
) -> None:
    client, factory, storage = processing_api_environment
    storage_key = add_document(
        factory,
        storage,
        content=b"invalid utf-8: \xff\xfe",
    )

    response = client.post(
        "/documents/1/process",
        headers=authorization_header(42),
    )

    assert response.status_code == 422
    body = response.text
    assert storage_key not in body
    assert str(storage._root) not in body  # noqa: SLF001 - security leak assertion
    assert "traceback" not in body.lower()
    assert "utf-8" not in body.lower()
