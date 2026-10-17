"""Generate neutral DOCX fixtures used to audit the browser preview renderer."""

from __future__ import annotations

import base64
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT_DIR = Path(__file__).with_name("docx")
VIETNAMESE_TEXT = "Tiếng Việt rõ ràng: học tập, kiến thức, kỹ năng và trách nhiệm."
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M/wHwAF"
    "gAJ/lCWaWQAAAABJRU5ErkJggg=="
)


def set_run_fonts(run, *, ascii_font=None, hansi_font=None, east_asia_font=None):
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    if ascii_font:
        fonts.set(qn("w:ascii"), ascii_font)
    if hansi_font:
        fonts.set(qn("w:hAnsi"), hansi_font)
    if east_asia_font:
        fonts.set(qn("w:eastAsia"), east_asia_font)


def set_character_spacing(run, value: int):
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(value))
    run._element.get_or_add_rPr().append(spacing)


def add_font_paragraph(document, font_name: str, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_fonts(
        run,
        ascii_font=font_name,
        hansi_font=font_name,
        east_asia_font=font_name,
    )
    run.font.size = Pt(12)
    return paragraph


def create_typography_fixture(path: Path):
    document = Document()
    document.add_heading("Kiểm thử font và tiếng Việt", level=1)
    for font in ("Times New Roman", "Arial", "Calibri", "Aptos"):
        add_font_paragraph(document, font, f"{font}: {VIETNAMESE_TEXT}")

    mixed = document.add_paragraph()
    for font, text in (
        ("Times New Roman", "Times New Roman — "),
        ("Arial", "Arial — "),
        ("Calibri", "Calibri — "),
        ("Aptos", "Aptos"),
    ):
        run = mixed.add_run(text)
        set_run_fonts(run, ascii_font=font, hansi_font=font)

    hansi_only = document.add_paragraph()
    run = hansi_only.add_run("Chỉ khai báo hAnsi: " + VIETNAMESE_TEXT)
    # Deliberately omit w:ascii. docx-preview 0.4.1 ignores an hAnsi-only
    # declaration, so Arial makes the fallback behavior observable in tests.
    set_run_fonts(run, hansi_font="Arial")

    spaced = document.add_paragraph()
    run = spaced.add_run("Character spacing được khai báo trong OOXML")
    set_run_fonts(run, ascii_font="Arial", hansi_font="Arial")
    set_character_spacing(run, 40)

    emphasis = document.add_paragraph()
    emphasis.add_run("Đậm").bold = True
    emphasis.add_run(" — ")
    emphasis.add_run("Nghiêng").italic = True
    emphasis.add_run(" — ")
    emphasis.add_run("Gạch chân").underline = True
    document.save(path)


def create_tables_fixture(path: Path):
    document = Document()
    document.add_heading("Kiểm thử bảng", level=1)

    left_table = document.add_table(rows=3, cols=4)
    left_table.style = "Table Grid"
    for row_index, row in enumerate(left_table.rows):
        for col_index, cell in enumerate(row.cells):
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.add_run(
                f"Ô {row_index + 1}.{col_index + 1}: Nội dung tiếng Việt căn trái."
            )

    document.add_paragraph()
    justified_table = document.add_table(rows=3, cols=3)
    justified_table.style = "Table Grid"
    long_text = (
        "Nội dung dài trong ô bảng được căn đều để so sánh cách dàn khoảng trắng "
        "giữa trình duyệt và Microsoft Word."
    )
    for row in justified_table.rows:
        for cell in row.cells:
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.add_run(long_text)
    document.save(path)


def create_lists_fixture(path: Path):
    document = Document()
    document.add_heading("Kiểm thử danh sách và thụt lề", level=1)
    for text in ("Mục đánh số thứ nhất", "Mục đánh số thứ hai"):
        document.add_paragraph(text, style="List Number")
    for text in ("Mục dấu đầu dòng thứ nhất", "Mục dấu đầu dòng thứ hai"):
        document.add_paragraph(text, style="List Bullet")

    hanging = document.add_paragraph(
        "Đoạn văn hanging indent dùng để xác nhận ký tự đầu dòng không bị cắt và "
        "nội dung vẫn giữ đúng thứ tự."
    )
    hanging.paragraph_format.left_indent = Inches(0.6)
    hanging.paragraph_format.first_line_indent = Inches(-0.3)
    document.save(path)


def create_multipage_media_fixture(path: Path):
    document = Document()
    image_path = path.with_suffix(".png")
    image_path.write_bytes(TINY_PNG)
    try:
        for page in range(1, 4):
            document.add_heading(f"Trang kiểm thử {page}", level=1)
            document.add_paragraph(VIETNAMESE_TEXT * 4)
            if page == 1:
                document.add_picture(str(image_path), width=Inches(0.5))
            if page < 3:
                document.add_page_break()
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("Phần cuối của tài liệu nhiều trang.")
        document.save(path)
    finally:
        image_path.unlink(missing_ok=True)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    create_typography_fixture(OUTPUT_DIR / "typography-vietnamese.docx")
    create_tables_fixture(OUTPUT_DIR / "tables-left-justify.docx")
    create_lists_fixture(OUTPUT_DIR / "lists-indentation.docx")
    create_multipage_media_fixture(OUTPUT_DIR / "multipage-media.docx")


if __name__ == "__main__":
    main()
